#!/usr/bin/env python3
"""Offline timetable-window calculation. No HTTP, database or operating-config access."""
import argparse
import csv
import io
import math
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, replace
from datetime import time
from pathlib import Path

DAY = 24 * 60 * 60


class InputError(ValueError):
    pass


def read_csv(path):
    raw = Path(path).read_bytes()
    for encoding in ('utf-8-sig', 'cp949'):
        try:
            return list(csv.DictReader(io.StringIO(raw.decode(encoding))))
        except UnicodeDecodeError:
            pass
    raise InputError(f'Cannot decode CSV: {path}')


def canonical(value):
    return re.sub(r'^TBS', 'TSB', str(value).strip().upper())


def seconds(value):
    if not re.fullmatch(r'\d{1,2}:\d{2}',str(value)):
        raise InputError(f'Invalid clock: {value}')
    h,m = map(int,value.split(':'))
    if h > 24 or m > 59 or h==24 and m:
        raise InputError(f'Invalid clock: {value}')
    return h*3600+m*60


def clock(value):
    if value==DAY:
        return '24:00'
    return f'{value//3600:02d}:{value%3600//60:02d}'


class Rules:
    def __init__(self, values):
        self.values = values
    @classmethod
    def load(cls,path):
        rows=read_csv(path)
        if any(not {'key','value'} <= set(r) for r in rows):
            raise InputError('Rules require key,value columns')
        if len({r['key'] for r in rows}) != len(rows):
            raise InputError('Duplicate window rule')
        return cls({r['key']:r['value'].strip() for r in rows})
    def number(self,key):
        try:
            value=float(self.values[key])
            if value < 0 or not math.isfinite(value):
                raise ValueError()
            return value
        except (KeyError,ValueError) as error:
            raise InputError(f'Missing/invalid numeric rule: {key}') from error
    def integers(self,key):
        try:
            values=[int(v) for v in self.values[key].split('|')]
            if not values or any(v <= 0 for v in values):
                raise ValueError()
            return sorted(set(values))
        except (KeyError,ValueError) as error:
            raise InputError(f'Invalid interval rule: {key}') from error


@dataclass(frozen=True)
class Departure:
    sheet_row: int
    side: str
    departure: int
    route_name: str
    trip_seconds: int
    keycode: str = ''
    route_id: str = ''
    duration_source: str = ''


def parse_departures(rows,rules):
    departures,warnings,audits=[],[],[]
    def get(row,key):
        column=int(rules.number(key))
        return row[column-1] if 0 < column <= len(row) else None
    for row_number,row in enumerate(rows,1):
        if row_number < rules.number('first_data_row'):
            continue
        label=get(row,'route_column')
        if label is None or not re.fullmatch(r'\d+(?:-\d+)?',str(label)):
            continue
        label=str(label)
        explicit=get(row,'trip_minutes_column')
        distance=get(row,'distance_column')
        if explicit is not None:
            if not isinstance(explicit,(int,float)) or explicit<=0:
                raise InputError(f'Invalid trip minutes at row {row_number}')
            duration=math.ceil(explicit*60); duration_source='시트 소요시간'
        elif rules.values.get('distance_speed_kmh') and isinstance(distance,(int,float)) and distance>0:
            speed=rules.number('distance_speed_kmh')
            if speed<=0:
                raise InputError('distance_speed_kmh must be positive')
            duration=math.ceil(distance/speed*3600); duration_source='거리/설정 속도'
        else:
            duration=math.ceil(rules.number('default_trip_minutes')*60); duration_source='기본 소요시간'
        if duration<=0:
            raise InputError('Trip duration must be positive')
        key=get(row,'keycode_column')
        key=str(key).strip() if key is not None else ''
        before=len(departures)
        for side,column in (('기점','origin_departure_column'),('종점','destination_departure_column')):
            raw=get(row,column)
            if raw is None or str(raw).strip()=='':
                warnings.append(f'시각 공란: 행 {row_number} 노선 {label} {side}; 출발 생성 안 함')
                continue
            if isinstance(raw,time):
                times=[raw.hour*3600+raw.minute*60+raw.second]
            elif isinstance(raw,(int,float)) and 0 <= raw < 1:
                times=[round(raw*DAY)]
            else:
                tokens=re.findall(r'(?<!\d)(\d{1,2}:\d{2})(?!\d)',str(raw))
                times=[seconds(v) for v in tokens]
                if not times:
                    warnings.append(f'시각 없는 메모: 행 {row_number} 노선 {label} {side}: {raw!s}')
            previous=-1
            for departure in times:
                if departure>=DAY:
                    raise InputError(f'24:00 departure requires explicit service-date handling: row {row_number}')
                if departure<previous:
                    warnings.append(f'시각 역전: 행 {row_number} 노선 {label} {side} {clock(departure)}; 원본 유지')
                previous=departure
                departures.append(Departure(row_number,side,departure,label,duration,key,duration_source=duration_source))
        declared=get(row,'declared_count_column')
        actual=len(departures)-before
        if isinstance(declared,(int,float)) and declared!=actual:
            audits.append(f'행 {row_number} 노선 {label}: F열 {declared:g}회, 실제 기점/종점 시각 {actual}회')
    if not departures:
        raise InputError('No timetable departures found')
    return departures,warnings,audits


def read_workbook(path,sheet,rules):
    if not Path(path).is_file():
        raise InputError(f'Required timetable not found: {path}')
    try:
        import openpyxl
        workbook=openpyxl.load_workbook(path,read_only=True,data_only=True)
        try:
            if sheet not in workbook.sheetnames:
                raise InputError(f'Required sheet not found: {sheet}')
            return parse_departures(workbook[sheet].iter_rows(values_only=True),rules)
        finally:
            workbook.close()
    except InputError:
        raise
    except Exception as error:
        raise InputError(f'Cannot read required timetable: {path}: {error}') from error


def load_mapping(path,rules):
    keys,names,routes,hundred=defaultdict(set),defaultdict(set),{},set()
    for row in read_csv(path):
        if not {'ROUTE_ID','ROUTE_NM','OBE_NM'} <= set(row):
            raise InputError('Mapping requires ROUTE_ID,ROUTE_NM,OBE_NM')
        route=canonical(row['ROUTE_ID']); name=row['ROUTE_NM'].strip(); key=row['OBE_NM'].strip()
        if not re.fullmatch(r'TSB\d{9}',route):
            raise InputError(f'Invalid mapped ROUTE_ID: {route}')
        routes[route]=name
        keys[key].add(route); names[name].add(route)
        if name.endswith('100'):
            hundred.add(route)
    if len(hundred)!=rules.number('expected_hundred_routes'):
        raise InputError(f'100번 집계 수 불일치: {len(hundred)}: {sorted(hundred)}')
    local=set(routes)-hundred
    if len(local)!=rules.number('expected_local_routes'):
        raise InputError(f'관내 집계 수 불일치: {len(local)}: {sorted(local)}')
    return keys,names,routes,hundred,local


def load_departure_keys(path):
    overrides={}
    if not path:
        return overrides
    for row in read_csv(path):
        if not {'sheet_row','side','departure_time','keycode'} <= set(row):
            raise InputError('Departure-key CSV requires sheet_row,side,departure_time,keycode')
        key=(int(row['sheet_row']),row['side'].strip(),seconds(row['departure_time']))
        if key in overrides:
            raise InputError(f'Duplicate departure-key override: {key}')
        overrides[key]=row['keycode'].strip()
    return overrides


def load_pattern_trips(path,mapping):
    if not path:
        return set()
    keys,names,routes,hundred,local=mapping
    raw=Path(path).read_bytes()
    try:
        text=raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text=raw.decode('cp949')
    trips=set()
    errors=[]
    for row in csv.reader(io.StringIO(text)):
        indices=[i for i,v in enumerate(row[:2]) if re.fullmatch(r'(?:ud|ld)\d+',v.strip())]
        if not indices:
            continue
        i=indices[0]
        for j in range(i+1,len(row),2):
            at=row[j].strip(); token=row[j+1].strip() if j+1<len(row) else ''
            if not at or not token:
                continue
            route=canonical(token)
            candidates={route} if route in routes else keys.get(token,set())
            if len(candidates)!=1:
                errors.append(f'{row[i]} {at} {token}: {sorted(candidates)}')
                continue
            route=next(iter(candidates))
            if route in local:
                trips.add((route,seconds(at)))
    if errors:
        raise InputError('Pattern keycode mapping failures:\n'+'\n'.join(errors))
    return trips


def map_departures(departures,mapping,overrides,pattern_trips=frozenset()):
    keys,names,routes,hundred,local=mapping
    mapped,failures=[],[]
    valid_keys={(d.sheet_row,d.side,d.departure) for d in departures}
    if set(overrides)-valid_keys:
        raise InputError(f'Departure-key rows absent from timetable: {sorted(set(overrides)-valid_keys)}')
    for d in departures:
        key=overrides.get((d.sheet_row,d.side,d.departure),d.keycode)
        if key:
            candidates=keys.get(key,set())
        else:
            candidates=names.get(d.route_name,set())
            if len(candidates)!=1:
                # The user confirmed these pattern ROUTE_IDs were mapped from TS OBE_NM.
                # Use only an exact departure-time match within the same named route family.
                # Never select the first route or an arbitrary nearest time/return trip.
                family_candidates=names.get(d.route_name.split('-')[0],set())
                exact={r for r,at in pattern_trips if r in family_candidates and at==d.departure}
                if exact:
                    candidates=exact
        if len(candidates)!=1:
            failures.append(dict(sheet_row=d.sheet_row,side=d.side,departure_time=clock(d.departure),
                                 route_name=d.route_name,keycode=key,candidate_route_ids='|'.join(sorted(candidates)),
                                 reason='keycode/노선명 매핑 없음' if not candidates else '복수 ROUTE_ID: keycode 확인 필요'))
            continue
        route=next(iter(candidates))
        if route not in local:
            failures.append(dict(sheet_row=d.sheet_row,side=d.side,departure_time=clock(d.departure),
                                 route_name=d.route_name,keycode=key,candidate_route_ids=route,
                                 reason='100번 계열은 관내 창에 포함할 수 없음'))
            continue
        mapped.append(replace(d,route_id=route,keycode=key))
    return mapped,failures


def split_midnight(start,end):
    if end<=start or end-start>DAY:
        raise InputError('Window must be positive and not longer than one day')
    result=[]
    while start<end:
        day=math.floor(start/DAY)
        boundary=min(end,(day+1)*DAY)
        result.append((start-day*DAY,boundary-day*DAY))
        start=boundary
    return result


def build_windows(departures,rules):
    grouped=defaultdict(list)
    for d in departures:
        if not d.route_id:
            raise InputError('Unresolved departure cannot create a window')
        for start,end in split_midnight(d.departure-round(rules.number('pre_departure_minutes')*60),
                                        d.departure+d.trip_seconds+round(rules.number('post_trip_minutes')*60)):
            grouped[d.route_id].append((start,end,{clock(d.departure)}))
    result={}
    for route,windows in grouped.items():
        merged=[]
        for start,end,departures_in_window in sorted(windows,key=lambda w:(w[0],w[1])):
            if merged and start<=merged[-1][1]:
                old_start,old_end,old_departures=merged[-1]
                merged[-1]=(old_start,max(old_end,end),old_departures|departures_in_window)
            else:
                merged.append((start,end,set(departures_in_window)))
        result[route]=merged
    return result


def calls(start,end,interval):
    if interval<=0 or end<start:
        raise InputError('Invalid call interval/window')
    return math.ceil((end-start)/interval)


def budget(local_calls,hundred_count,rules):
    hundred=hundred_count*calls(seconds(rules.values['hundred_start']),seconds(rules.values['hundred_end']),rules.number('hundred_interval_seconds'))
    result=[]
    for local_interval,local in sorted(local_calls.items()):
        for backup_interval in rules.integers('backup_intervals_seconds'):
            backup=int(rules.number('backup_route_count'))*calls(seconds(rules.values['backup_start']),seconds(rules.values['backup_end']),backup_interval)
            total=local+hundred+backup
            remaining=int(rules.number('daily_budget'))-total
            result.append(dict(local_interval_seconds=local_interval,backup_interval_seconds=backup_interval,
                               local_calls=local,hundred_calls=hundred,backup_calls=backup,total_calls=total,
                               budget_limit=int(rules.number('daily_budget')),remaining_calls=remaining,
                               within_budget=remaining>=0,retry_reserve_required=int(rules.number('retry_reserve')),
                               retry_reserve_ok=remaining>=rules.number('retry_reserve')))
    return result


def cross_validate(path,departures,mapping):
    if not path:
        return []
    keys,names,routes,hundred,local=mapping
    expected={(d.route_id,d.departure) for d in departures}
    differences=[]
    raw=Path(path).read_bytes()
    try:
        text=raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text=raw.decode('cp949')
    for row in csv.reader(io.StringIO(text)):
        indices=[i for i,v in enumerate(row[:2]) if re.fullmatch(r'(?:ud|ld)\d+',v.strip())]
        if not indices:
            continue
        i=indices[0]; pattern=row[i]
        for j in range(i+1,len(row),2):
            at=row[j].strip(); token=row[j+1].strip() if j+1<len(row) else ''
            if not at and not token:
                continue
            if not at or not token:
                differences.append(dict(pattern=pattern,time=at,route_id=token,reason='패턴 시각/노선 공란; 시간표 우선'))
                continue
            route=canonical(token)
            candidates={route} if route in routes else keys.get(token,set())
            if len(candidates)!=1:
                raise InputError(f'Cross-validation pattern mapping failed: {pattern} {at} {token}')
            route=next(iter(candidates))
            if route not in local or (route,seconds(at)) not in expected:
                differences.append(dict(pattern=pattern,time=at,route_id=route,reason='시간표와 다름; 시간표 우선'))
    return differences


def write_csv(path,rows,fields):
    with Path(path).open('w',encoding='utf-8-sig',newline='') as out:
        writer=csv.DictWriter(out,fieldnames=fields); writer.writeheader(); writer.writerows(rows)


def sequence_score(reference,candidate):
    """LCS/reference length, exact normalized names; no aliases or direction inference."""
    if not reference or not candidate:
        return ''
    previous=[0]*(len(candidate)+1)
    for value in reference:
        current=[0]
        for j,other in enumerate(candidate,1):
            current.append(previous[j-1]+1 if value==other else max(previous[j],current[-1]))
        previous=current
    return round(100*previous[-1]/len(reference),2)


def unresolved_details(failures,mapping,workbook_path,sheet,stops_path=None):
    keys,names,routes,hundred,local=mapping
    references={}
    candidates_stops=defaultdict(list)
    if stops_path:
        import openpyxl
        book=openpyxl.load_workbook(workbook_path,read_only=True,data_only=True)
        try:
            anchor=None
            for i,row in enumerate(book[sheet].iter_rows(values_only=True),1):
                if row and row[0] is not None and re.fullmatch(r'\d+(?:-\d+)?',str(row[0])):
                    anchor=i; references[anchor]=[]
                if anchor is not None:
                    references[anchor].extend(re.sub(r'\s+','',str(v)) for v in row[10:] if v is not None)
        finally:
            book.close()
        stop_rows=read_csv(stops_path)
        required={'route_id','api_response_sequence','node_name'}
        if any(not required <= set(r) for r in stop_rows):
            raise InputError('Route stops require route_id,api_response_sequence,node_name')
        for r in sorted(stop_rows,key=lambda r:(canonical(r['route_id']),int(r['api_response_sequence']))):
            candidates_stops[canonical(r['route_id'])].append(re.sub(r'\s+','',r['node_name']))
    result=[]
    for f in failures:
        candidates=set(filter(None,f['candidate_route_ids'].split('|')))
        if not candidates:
            candidates=names.get(f['route_name'].split('-')[0],set()) & local
        candidate_keys=sorted(k for k,ids in keys.items() if ids & candidates)
        scores=[]
        for key in candidate_keys:
            for route in sorted(keys[key] & candidates):
                score=sequence_score(references.get(f['sheet_row'],[]),candidates_stops.get(route,[]))
                scores.append(f'{key}:{route}={score}' if score!='' else f'{key}:{route}=')
        result.append(dict(sheet_row=f['sheet_row'],route_name=f['route_name'],departure_time=f['departure_time'],
                           direction=f['side'],candidate_keycodes='|'.join(candidate_keys),
                           candidate_stop_order_match_pct='|'.join(scores),collection_mode='백업 스윕 전용',
                           score_basis='시간표 전체 경유지 대비 정규화 이름의 순서 일치(LCS); 방향 추정 없음' if stops_path else '정류장 순서 자료 미지정',
                           reason=f['reason']))
    return result


def concurrency(windows):
    boundaries={0,DAY}
    for items in windows.values():
        for start,end,_ in items:
            boundaries.update((start,end))
    points=sorted(boundaries)
    return [dict(time=clock(start),until=clock(end),active_route_count=len(active),route_ids='|'.join(active))
            for start,end in zip(points,points[1:])
            for active in [sorted(route for route,items in windows.items() if any(a<=start<b for a,b,_ in items))]]


def confirmed_mapping_overrides(path,failures,mapping):
    if not path:
        return {}
    allowed={(f['route_name'],f['side'],f['departure_time']):f for f in failures}
    keys,names,routes,hundred,local=mapping
    result={}
    for row in read_csv(path):
        identity=(row['route_name'],row['direction'],row['departure_time'])
        failure=allowed.get(identity)
        key=row['keycode']; route=row['route_id']
        if (failure is None or row.get('grade')!='확정 후보' or keys.get(key)!= {route}
                or route not in local or routes[route].split('-')[0]!=row['route_name'].split('-')[0]
                or int(row['sheet_row'])!=int(failure['sheet_row'])):
            raise InputError(f'Invalid confirmed mapping: {identity}')
        target=(int(row['sheet_row']),row['direction'],seconds(row['departure_time']))
        if target in result:
            raise InputError(f'Duplicate confirmed mapping: {identity}')
        result[target]=key
    return result


def run(args):
    rules=Rules.load(args.rules)
    departures,warnings,audits=read_workbook(args.workbook,args.sheet,rules)
    mapping=load_mapping(args.mapping,rules)
    pattern_trips=load_pattern_trips(args.patterns,mapping)
    mapped,failures=map_departures(departures,mapping,load_departure_keys(args.departure_keys),pattern_trips)
    base_count=len(mapped)
    extra=confirmed_mapping_overrides(getattr(args,'confirmed_mapping',None),failures,mapping)
    if extra:
        overrides=load_departure_keys(args.departure_keys)
        overrides.update(extra)
        mapped,failures=map_departures(departures,mapping,overrides,pattern_trips)
        if len(mapped)!=base_count+len(extra):
            raise InputError('Confirmed mapping count mismatch')
    expected=int(rules.number('expected_departures'))
    if len(departures)!=expected or len(mapped)+len(failures)!=expected:
        raise InputError(f'Departure total mismatch: {len(mapped)} + {len(failures)} != {expected}')
    unresolved=unresolved_details(failures,mapping,args.workbook,args.sheet,getattr(args,'route_stops',None))
    out=Path(args.output_dir); out.mkdir(parents=True,exist_ok=True)
    lines=[f'시간표: {Path(args.workbook).resolve()} / {args.sheet}',
           f'관내 출발 총횟수: {len(departures)} / 기준 {int(rules.number("expected_departures"))}회: '+
           ('일치' if len(departures)==rules.number('expected_departures') else '불일치: 시간표에서 읽은 실제 시각을 유지'),
           f'관내 노선 범위: {len(mapping[4])}개, 100번 {len(mapping[3])}개 별도',
           '출발 조건: 운수사 시간표의 출발 기점/시각 기준. TAGO 노선 기점은 정류장 순번 최소값. 특정 정류장 관측 유무로 창 생성 여부를 결정하지 않음.',
           f'확정: {len(mapped)}회 ({len(mapped)/expected:.2%}); 미확정: {len(failures)}회 ({len(failures)/expected:.2%}); 합계 {expected}회',
           f'창 계산 입력: 기본 {base_count}회 + confirmed_mapping {len(extra)}회; 입력 파일: {getattr(args,"confirmed_mapping",None)}',
           f'확정 연결: {len(mapped)} / {len(departures)}회; 패턴의 TS 매핑 ROUTE_ID·정확한 시각 연결 사용',
           '소요시간: 명시적 소요값 우선, 거리 환산 속도 미지정 시 기본값 사용; 1일 운행거리는 소요시간이 아님.',
           '창은 [시작 포함, 종료 제외), 같은 노선만 병합, 자정에서 분할 후 일별 반복 창으로 정규화.',
           '백업 스윕은 설정한 시간대의 전체 노선 추가 호출로 계산(다른 호출과 상쇄하지 않음).',
           f'재시도 여유는 운영 예산 {int(rules.number("daily_budget"))} 안에 남는 호출 수 기준; 정책 값은 window_rules.csv에서 변경.',
           '[F열 운행횟수와 실제 시각 차이]']+audits+['[시각 역전/공란/메모 경고]']+warnings
    failure_fields=['sheet_row','side','departure_time','route_name','keycode','candidate_route_ids','reason']
    write_csv(out/'mapping_failures.csv',failures,failure_fields)
    write_csv(out/'unresolved_departures.csv',unresolved,['sheet_row','route_name','departure_time','direction',
              'candidate_keycodes','candidate_stop_order_match_pct','collection_mode','score_basis','reason'])
    windows=build_windows(mapped,rules)
    local_calls={interval:0 for interval in rules.integers('local_intervals_seconds')}
    plan,summaries=[],[]
    for route in sorted(mapping[4]):
        items=windows.get(route,[])
        confirmed=sum(d.route_id==route for d in mapped)
        summary=dict(route_id=route,route_name=mapping[2][route],departure_count=confirmed,
                     confirmed_departure_count=confirmed,unresolved_departure_count=0,
                     collection_mode='확정 창 + 백업 스윕' if confirmed else '백업 스윕 전용',window_count=len(items))
        for interval in local_calls:
            count=sum(calls(start,end,interval) for start,end,_ in items)
            summary[f'calls_{interval}s']=count; local_calls[interval]+=count
            for start,end,times in items:
                plan.append(dict(route_id=route,route_name=mapping[2][route],departure_time='|'.join(sorted(times)),
                                 window_start=clock(start),window_end=clock(end),interval_seconds=interval,calls_per_window=calls(start,end,interval)))
        summaries.append(summary)
    # Unresolved counts have no unique ROUTE_ID: do not allocate them to each candidate.
    for name in sorted({r['route_name'] for r in unresolved}):
        summaries.append(dict(route_id='',route_name=name,departure_count=0,confirmed_departure_count=0,
                              unresolved_departure_count=sum(r['route_name']==name for r in unresolved),
                              collection_mode='백업 스윕 전용',window_count=0,**{f'calls_{v}s':0 for v in local_calls}))
    concurrent=concurrency(windows)
    peak=max(r['active_route_count'] for r in concurrent)
    peak_times=[f'{r["time"]}~{r["until"]}' for r in concurrent if r['active_route_count']==peak]
    differences=cross_validate(args.patterns,mapped,mapping)
    budgets=budget(local_calls,len(mapping[3]),rules)
    write_csv(out/'window_plan.csv',plan,['route_id','route_name','departure_time','window_start','window_end','interval_seconds','calls_per_window'])
    write_csv(out/'window_summary.csv',summaries,['route_id','route_name','departure_count','confirmed_departure_count',
              'unresolved_departure_count','collection_mode','window_count']+[f'calls_{v}s' for v in local_calls])
    write_csv(out/'concurrency.csv',concurrent,['time','until','active_route_count','route_ids'])
    write_csv(out/'budget_check.csv',budgets,list(budgets[0]))
    write_csv(out/'timetable_pattern_differences.csv',differences,['pattern','time','route_id','reason'])
    lines+=['[노선별 적용 소요시간]']+sorted({f'{d.route_name}: {d.trip_seconds/60:g}분 ({d.duration_source})' for d in mapped})
    lines+=['상태: 확정분 계산 완료; 미확정분은 창 생성 없이 백업 스윕 전용',
            '미확정 ROUTE_ID는 공란 행에 원 시간표 노선명으로 집계; 후보별 중복 집계/자동 배정 없음.',
            '정류장 점수는 전체 경유지 기준 참고값이며 방향별 확정 근거가 아님. 이름 별칭 추정 없음.',
            f'관내 창 동시 활성 노선 최댓값: {peak}개; 발생 구간: '+', '.join(peak_times),
            '동시 활성 수는 관내 확정 창만 계산; 100번 상시/백업 스윕은 포함하지 않음.',
            f'패턴 교차검증 차이: {len(differences)}건 (시간표 우선)',f'하루 관내 호출 수: {local_calls}','[예산 판정]']
    lines+=[str(row) for row in budgets]
    (out/'summary.txt').write_text('\n'.join(lines)+'\n',encoding='utf-8-sig')
    print('\n'.join(lines))
    return 0


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',required=True)
    parser.add_argument('--sheet',default='관내')
    parser.add_argument('--mapping',required=True)
    parser.add_argument('--patterns')
    parser.add_argument('--route-stops',help='Optional offline route stop CSV; scores are diagnostic only')
    parser.add_argument('--departure-keys',help='Optional explicit sheet_row,side,departure_time,keycode CSV; no inferred variants')
    parser.add_argument('--confirmed-mapping',help='Aggregated single-candidate evidence CSV from unresolved_recheck.py')
    parser.add_argument('--rules',default=str(Path(__file__).parent/'config'/'window_rules.csv'))
    parser.add_argument('--output-dir',default='reports/window-plan')
    args=parser.parse_args()
    try:
        return run(args)
    except (InputError,OSError,ValueError) as error:
        print(str(error),file=sys.stderr); return 2


if __name__=='__main__':
    raise SystemExit(main())
